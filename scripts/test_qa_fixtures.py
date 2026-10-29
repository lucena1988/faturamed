import io
import unittest
from openpyxl import load_workbook
from scripts.qa_fixtures import workbook_bytes


class FixtureTest(unittest.TestCase):
    def test_workbook_without_existing_template_preserves_literal_data(self):
        rows = [['Atendimento', 'Data Consumo', 'Medico'],
                ['000123', '2026-06-10', 'Medico & <QA>'],
                ['000124', '2026-06-11', '=1+1']]
        workbook = load_workbook(io.BytesIO(workbook_bytes(rows)))
        self.assertEqual(workbook.sheetnames, ['Planilha2'])
        self.assertEqual(list(workbook.active.values), [tuple(row) for row in rows])
        self.assertEqual(workbook.active['C3'].data_type, 's')
        workbook.close()

    def test_wide_sheet_and_numeric_values(self):
        rows = [[str(i) for i in range(30)], [0, 150.5]]
        workbook = load_workbook(io.BytesIO(workbook_bytes(rows)))
        self.assertEqual(workbook.active['AD1'].value, '29')
        self.assertEqual(workbook.active['A2'].value, '0')
        self.assertEqual(workbook.active['B2'].value, '150.5')
        workbook.close()


if __name__ == '__main__':
    unittest.main()
